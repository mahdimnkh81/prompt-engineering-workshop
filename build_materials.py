from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from curriculum import TASKS, AGENDA
from copy import deepcopy
TASKS=deepcopy(TASKS)
for task in TASKS:
    task['brief']=task['ai_brief']
    task['artifact']='Only your prompt. The AI evaluator returns scores, evidence, and improvement advice.'

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'materials';OUT.mkdir(exist_ok=True)
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
NAVY='13243B';TEAL='46DDC2';WHITE='F6F8FC';MUTED='B9C8D9'
notes=[]
def box(slide,x,y,w,h,text,size=24,color=WHITE,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(16)
    return shape

def slide(title,body,kicker='PROMPT LAB',note='',size=26):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY)
    box(s,.7,.35,12,.45,kicker,12,TEAL,True)
    box(s,.7,1,12,1.25,title,36,WHITE,True)
    box(s,.7,2.45,11.85,4.2,body,size)
    box(s,.7,7,11,.3,'PROMPT ENGINEERING  /  PRACTICE · CHECK · REFINE',10,MUTED)
    box(s,12,7,.6,.3,str(len(prs.slides)),10,TEAL)
    s.notes_slide.notes_text_frame.text=note
    notes.append(f'## Slide {len(prs.slides)} — {title}\n\n{note}\n')

slide('Better prompts. Measurable results.','A three-hour workshop for 20 participants\nEight practical labs • individual workspaces • feedback at every step',note='00:00–00:03. Welcome participants. Explain that success means usable, verified outputs. Ask everyone to choose one real task they want to improve.')
slide('One model. Two very different briefs.','Vague: “Write something about our product.”\nSpecific: “Write a two-sentence LinkedIn ad for a project-management tool. Audience: operations managers. Confident tone. End with a free-trial CTA. No emojis.”',note='00:03–00:07. Source p1. Ask which assumptions the vague request leaves open. Compare role, audience, output shape, and constraints; do not promise a universal improvement percentage.',size=25)
slide('Your route through the workshop','00:00–00:50  Foundations + few-shot examples\n00:50–01:30  Verification + structured output\n01:30–01:40  Break\n01:40–02:30  Revision + interviews + workflows\n02:30–03:00  Capstone + transfer to your own work',note='00:07–00:08. Display schedule. Labs include explanation, practice, review and retry time. Review submissions continuously during practice.',size=25)
slide('Open your own workspace','Sign in with your assigned participant account.\nSubmit your prompt and answer the concept question.\nAI assessment: four weighted criteria, 100 points total.\nScore ≥70/100 and answer correctly to unlock the next task.',note='00:08–00:10. Confirm AvalAI is configured. Give each participant only their own credential row. Demonstrate Save draft, Submit, and Refresh feedback. Failed API requests do not assign a score or unlock a task.',size=25)

for t in TASKS:
    if t['id']==5:
        slide('Break · 10 minutes','Return at 01:40.\nLeave your draft saved.\nInstructor: help participants improve prompts that remain below 70.',note='01:30–01:40. Preserve the break. Use the dashboard to identify participants needing help.')
    slot=next(a[0] for a in AGENDA if a[2].startswith(str(t['id'])+' ·'))
    teach=4 if t['minutes']==20 else 3
    practice=10 if t['minutes']==20 else 7
    review=t['minutes']-teach-practice
    slide(t['title'],t['concept']+'\n\n'+('Ask for concise calculation summaries and verify independently.' if t['id']==3 else 'Make the expected behavior observable before testing.'),kicker=f"LAB {t['id']:02}  /  {slot}  /  SOURCE PAGES {t['pages']}",note=f"First {teach} minutes of {slot}: explain the concept, then demonstrate the next slide. Source PDF pages {t['pages']}. Ask a participant to name one likely failure mode.",size=27)
    if t['id']==1:
        slide('Demonstration: explain an API invoice item','You are developing software for a nontechnical client. Explain what an API is and why this work appears on the invoice. Write a prompt using role, audience, tone, and format.',kicker='LAB 01 / TEACHING EXAMPLE',note='Use this API explanation as the demonstration. The participant exercise is the separate LinkedIn Python hiring scenario.',size=25)
    slide('A pattern you can adapt',t['example'],kicker=f"LAB {t['id']:02}  /  DEMONSTRATION",note='Demonstrate this example as part of the teaching allocation. Ask which details are reusable and which must change for a new task. Model outputs can vary; compare against requirements.',size=25)
    slide('Your task',t['brief'],kicker=f"LAB {t['id']:02} / {practice} MINUTES OF PRACTICE",note=f"Practice for {practice} minutes. Submit your prompt and the concept answer on your own page. The AI evaluator provides scores and feedback.",size=24)
    slide('Check, discuss, improve','\n'.join('• '+x for x in t['ai_rubric']),kicker=f"LAB {t['id']:02}  /  {review} MINUTES OF REVIEW & RETRY",note=f"Reserve {review} minutes for feedback, retries, and a short discussion. Use the task-specific maximum points: Lab 1 uses 20/20/20/40; other labs use 25 each. Ask: {t['quiz']} Correct answer: {t['options'][t['answer']]}. Never use model self-confidence as proof of correctness.",size=25)

slide('Common failures → useful repairs','Vague task → define audience, purpose, and output.\nOverloaded request → split into stages with clear handoffs.\nMissing context → include facts and examples.\nUnusable output → validate format and factual content.\nAssumed memory → restate the relevant context.',note='02:50–02:53. Source pp7–8. Ask participants which failure they observed in their own attempts.',size=25)
slide('Transfer the technique to your work','Writing: audience + purpose + CTA + length.\nCode: relevant code + error + one focused change + tests.\nData: source facts + schema + unknown-value policy.\nResearch: alternatives + evidence + uncertainty.',note='02:53–02:56. Source p8. Invite two participants to show a before/after example. Model-generated citations and APIs need independent verification.',size=27)
slide('Keep a prompt journal','Save the prompt, model/settings, test cases, and results.\nRecord one failed case and the change that helped.\nNext week: reuse one pattern on a recurring task.\nDownload your journal from your workspace.',note='02:56–03:00. Source p9. Ask for one specific next action. Point to source PDF and provider-specific documentation; controls such as temperature differ between models.',size=27)
prs.save(OUT/'prompt_engineering_3_hour_workshop.pptx')
(OUT/'speaker_notes.md').write_text('# Speaker notes\n\n'+ '\n'.join(notes))
agenda='\n'.join(f'| {time} | {mins} | {title} |' for time,mins,title in AGENDA)
(OUT/'facilitator_guide.md').write_text("""# Facilitator guide — AI prompt assessment

20 participants · 180 minutes · source: original nine-page PDF.

| Time | Minutes | Activity |
|---|---:|---|
"""+agenda+"""

## Before class

Configure AVALAI_API_KEY in the server environment or .streamlit/secrets.toml; see README.md. The configured provider is AvalAI, model gpt-6-astra, endpoint https://api.avalai.ir/v1. Test a real submission before class and confirm model access and sufficient provider credit. Give each participant only their individual credential row. Open the PowerPoint in presenter mode.

## Participant flow

Participants write a prompt and answer the original three-option concept question. They do not use an external chat model, paste a generated answer, or provide a reflection. The server sends their prompt as untrusted data to a separate evaluator instruction containing the task-specific rubric. The evaluator assesses the prompt without executing it and provides English feedback, evidence, issues, and improvements.

## Scoring

Lab 1: role 20, audience 20, tone 20, format 40 (structure 20 + explicit example/analogy 20). Missing the example/analogy request caps format at 20/40, not automatic failure. Other labs: four criteria × 25 = 100. Anchors: 0 absent/contradictory, 5 severe gaps, 10 partial, 15 mostly specified with important gaps, 20 clear with minor gaps, 25 complete. The server validates the response and sums the criterion scores itself. At least 70 plus a correct concept answer unlocks the next task immediately; otherwise participants must retry. The concept question is checked by the server and adds no points to the prompt score. API failures assign no score and preserve the draft. Existing earned passes and historical manual attempts remain available.

Model grading is an assessment aid, not proof of real output reliability. It can vary or misjudge a prompt. Review a sample of scores before class, discuss disputed feedback, and help participants revise. New submissions no longer require instructor approval. Historical attempts remain available in the dashboard history.

## Timing and facilitation

Use each lab's teaching, practice, and retry allocations in the slide notes. During practice, watch class progress and help participants who remain blocked. Strict mastery gates may mean some need follow-up after the three-hour session. Keep the scheduled break.

## Source mapping

Labs 1–8 map to PDF pages 1–3, 3, 3–4, 4, 5, 5–6, 6–7, and 7–9. The source's unsupported 80% claim is omitted. Verification is taught through concise checkable calculations. Participant code is never executed by the server.
""")
handout=['# Participant lab book\n\nUse your own workspace. Submit your prompt and answer the three-option concept question. The AI evaluator scores four weighted criteria totaling 100; at least 70/100 plus a correct concept answer unlocks the next lab. You may retry.\n']
for t in TASKS:
    handout.append(f"## Lab {t['id']}: {t['title']} ({t['minutes']} minutes)\n\n{t['concept']}\n\n**Task.** {t['brief']}\n\n**Submit.** {t['artifact']}\n\n**After feedback.** Read the scores, evidence, and suggestions. Revise your prompt and retry when below 70.\n\n**AI rubric (task-specific weights):**\n\n"+'\n'.join('- '+x for x in t['ai_rubric'])+'\n')
(OUT/'participant_lab_book.md').write_text('\n'.join(handout))
print(f'Created {len(prs.slides)} slides and workshop handouts.')
