"""
Remove Lab 8 concept and exercise slides from Prompt_Engineering_Workshop_v2.pptx and v3.pptx.
Also update the agenda slide to replace Lab 8 with 'Open Practice & Live Q&A'.
"""
from pathlib import Path
from pptx import Presentation

def remove_lab8(path):
    print(f"Processing {path}...")
    prs = Presentation(path)

    # 1. Update Agenda slide (Slide index 6)
    if len(prs.slides) > 6:
        agenda = prs.slides[6]
        for shp in agenda.shapes:
            if shp.has_text_frame:
                txt = shp.text_frame.text
                if 'Lab 8' in txt and 'Capstone' in txt:
                    shp.text_frame.text = 'Open Practice & Live Q&A'
                    for p in shp.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.name = 'Segoe UI'
                elif txt.strip() == '8':
                    shp.text_frame.text = 'Q&A'
                    for p in shp.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.name = 'Segoe UI'

    # 2. Find and delete Lab 8 slides
    # Look for slides with 'Lab 8' or 'Exercise 8'
    indices_to_remove = []
    for i, s in enumerate(prs.slides):
        if i == 6:  # skip agenda slide
            continue
        is_lab8 = False
        for shp in s.shapes:
            if shp.has_text_frame:
                t = shp.text_frame.text
                if 'Lab 8' in t or 'Exercise 8' in t or 'Capstone: Evidence' in t:
                    is_lab8 = True
                    break
        if is_lab8:
            indices_to_remove.append(i)

    print(f"Found Lab 8 slides at indices: {indices_to_remove}")

    # Remove in reverse order so indices do not shift
    xml_slides = prs.slides._sldIdLst
    slides_list = list(xml_slides)
    for idx in sorted(indices_to_remove, reverse=True):
        print(f"Removing slide at index {idx}...")
        xml_slides.remove(slides_list[idx])

    prs.save(path)
    print(f"Saved {path} successfully. Total remaining slides: {len(prs.slides)}.")

if __name__ == '__main__':
    v2_path = Path(r"d:\My work\workshop\prompt-engineering-workshop\materials\Prompt_Engineering_Workshop_v2.pptx")
    v3_path = Path(r"d:\My work\workshop\prompt-engineering-workshop\materials\Prompt_Engineering_Workshop_v3.pptx")

    if v2_path.exists():
        remove_lab8(v2_path)
    if v3_path.exists():
        remove_lab8(v3_path)
